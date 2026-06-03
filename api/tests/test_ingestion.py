"""Testes de roteamento da ingestão multimodal (``build_documents_from_file``).

Os modelos pesados (Gemini / Whisper) são substituídos por *fakes*; bibliotecas
de sistema (pdf2image/poppler) são monkeypatchadas. Validamos o roteamento por
extensão, a paginação e os metadados — sem baixar/rodar modelos reais.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from app.core.ingestion import (
    EmptyDocumentError,
    UnsupportedFileTypeError,
    build_documents_from_file,
)
from app.core.transcription import WhisperEngine, _format_timestamp


class FakeGemma:
    """Engine de visão falsa: registra chamadas e devolve um texto fixo por imagem."""

    def __init__(self, reply: str = "TEXTO OCR") -> None:
        self.reply = reply
        self.calls: list[tuple[str, str | None]] = []

    def describe_image(self, path: str, *, prompt: str | None = None) -> str:
        self.calls.append((path, prompt))
        return self.reply


class FakeWhisper:
    """Engine STT falsa: devolve uma transcrição com timestamp."""

    def __init__(self, reply: str = "[00:00:51] trecho de teste") -> None:
        self.reply = reply
        self.calls: list[str] = []

    def transcribe(self, path: str) -> str:
        self.calls.append(path)
        return self.reply


def _build(path, gemma=None, stt=None, **kw):
    return build_documents_from_file(
        path,
        engine=gemma or FakeGemma(),
        stt_engine=stt or FakeWhisper(),
        doc_id=kw.pop("doc_id", "doc"),
        metadata=kw.pop("metadata", None),
    )


# --------------------------------------------------------------------------- #
# Texto direto
# --------------------------------------------------------------------------- #
def test_txt_is_read_directly(tmp_path: Path) -> None:
    f = tmp_path / "report.txt"
    f.write_text("CWE-862 e CWE-284 encontrados.", encoding="utf-8")

    docs = _build(f, doc_id="rep")

    assert len(docs) == 1
    assert "CWE-862" in docs[0].text
    assert docs[0].metadata["page"] == 1
    assert docs[0].metadata["modality"] == "text"
    assert docs[0].metadata["source"] == "report.txt"


def test_content_hash_kept_in_metadata_but_excluded_from_embedding(tmp_path: Path) -> None:
    """O ``content_hash`` deve ser gravado nos metadados, mas NÃO poluir o embedding."""
    from llama_index.core.schema import MetadataMode

    f = tmp_path / "report.txt"
    f.write_text("conteudo semantico relevante", encoding="utf-8")

    docs = _build(f, doc_id="rep", metadata={"content_hash": "deadbeef" * 8})
    doc = docs[0]

    # Continua disponível para listar/excluir/atualizar...
    assert doc.metadata["content_hash"] == "deadbeef" * 8
    # ...mas fora do texto usado para embedding e para o LLM.
    assert "deadbeef" not in doc.get_content(metadata_mode=MetadataMode.EMBED)
    assert "deadbeef" not in doc.get_content(metadata_mode=MetadataMode.LLM)
    assert "conteudo semantico relevante" in doc.get_content(metadata_mode=MetadataMode.EMBED)


def test_unsupported_extension_raises(tmp_path: Path) -> None:
    f = tmp_path / "x.exe"
    f.write_bytes(b"\x00")
    with pytest.raises(UnsupportedFileTypeError):
        _build(f)


def test_empty_text_raises(tmp_path: Path) -> None:
    f = tmp_path / "empty.txt"
    f.write_text("   \n", encoding="utf-8")
    with pytest.raises(EmptyDocumentError):
        _build(f)


# --------------------------------------------------------------------------- #
# Áudio -> Whisper
# --------------------------------------------------------------------------- #
def test_audio_uses_whisper_engine() -> None:
    stt = FakeWhisper(reply="[00:00:51] Combo e venda casada continuam valendo.")
    docs = _build("alinhamento.mp3", stt=stt, doc_id="audio")

    assert stt.calls == ["alinhamento.mp3"]
    assert len(docs) == 1
    assert "00:00:51" in docs[0].text
    assert docs[0].metadata["modality"] == "audio"
    assert docs[0].metadata["page"] == 1


# --------------------------------------------------------------------------- #
# Imagem -> Gemma
# --------------------------------------------------------------------------- #
def test_image_uses_gemma(tmp_path: Path) -> None:
    img = tmp_path / "foto.png"
    img.write_bytes(b"\x89PNG\r\n")  # conteúdo não importa: describe_image é fake
    gemma = FakeGemma(reply="Treino Livre 3, Classificacao, Q1, Q2 e Q3")
    docs = _build(img, gemma=gemma, doc_id="img")

    assert len(gemma.calls) == 1
    assert "Q1, Q2 e Q3" in docs[0].text
    assert docs[0].metadata["modality"] == "image"


# --------------------------------------------------------------------------- #
# PDF -> pdf2image + Gemma (1 documento por página)
# --------------------------------------------------------------------------- #
class _FakePage:
    def __init__(self, idx: int) -> None:
        self.idx = idx

    def save(self, path, format: str | None = None) -> None:  # noqa: A002
        Path(path).write_bytes(b"img")


def test_pdf_routes_each_page_through_gemma(tmp_path: Path, monkeypatch) -> None:
    pdf = tmp_path / "doc.pdf"
    pdf.write_bytes(b"%PDF-1.4")

    import pdf2image

    monkeypatch.setattr(
        pdf2image, "convert_from_path", lambda *a, **k: [_FakePage(1), _FakePage(2), _FakePage(3)]
    )

    replies = iter(["pagina 1", "pagina 2", "DO NOT CONTINUE YOUR TASK"])

    class SeqGemma(FakeGemma):
        def describe_image(self, path: str, *, prompt: str | None = None) -> str:
            self.calls.append((path, prompt))
            return next(replies)

    gemma = SeqGemma()
    docs = _build(pdf, gemma=gemma, doc_id="pdf")

    assert len(docs) == 3
    assert [d.metadata["page"] for d in docs] == [1, 2, 3]
    assert [d.doc_id for d in docs] == ["pdf:p1", "pdf:p2", "pdf:p3"]
    assert "DO NOT CONTINUE YOUR TASK" in docs[2].text
    assert docs[0].metadata["modality"] == "pdf"


def test_pdf_all_pages_empty_raises(tmp_path: Path, monkeypatch) -> None:
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    import pdf2image

    monkeypatch.setattr(pdf2image, "convert_from_path", lambda *a, **k: [_FakePage(1)])
    with pytest.raises(EmptyDocumentError):
        _build(pdf, gemma=FakeGemma(reply="   "))


# --------------------------------------------------------------------------- #
# XLSX -> pandas + openpyxl (1 documento por aba)
# --------------------------------------------------------------------------- #
def test_xlsx_extracts_sheet_text(tmp_path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Notas"
    ws.append(["Disciplina", "Falta"])
    ws.append(["Sistemas Operacionais Embarcados", "0,54"])
    xlsx = tmp_path / "notas.xlsx"
    wb.save(xlsx)

    docs = _build(xlsx, doc_id="notas")

    assert len(docs) == 1
    assert "Sistemas Operacionais Embarcados" in docs[0].text
    assert "0,54" in docs[0].text
    assert docs[0].metadata["modality"] == "spreadsheet"
    assert docs[0].metadata["sheet"] == "Notas"
    assert docs[0].metadata["page"] == 1
    assert docs[0].doc_id == "notas:s1"


# --------------------------------------------------------------------------- #
# DOCX -> MarkItDown + python-docx
# --------------------------------------------------------------------------- #
def test_docx_extracts_text(tmp_path: Path) -> None:
    from docx import Document as DocxDocument

    d = DocxDocument()
    d.add_paragraph("Minha stack: Next.js, TypeScript, React, Node.js e Supabase.")
    docx = tmp_path / "carta.docx"
    d.save(docx)

    docs = _build(docx, doc_id="carta")

    assert len(docs) == 1
    assert "Next.js, TypeScript, React, Node.js e Supabase" in docs[0].text
    assert docs[0].metadata["modality"] == "document"
    assert docs[0].metadata["page"] == 1


# --------------------------------------------------------------------------- #
# WhisperEngine: formatação de timestamps (sem carregar modelo)
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    "seconds,expected",
    [(0, "00:00:00"), (51, "00:00:51"), (75, "00:01:15"), (3661, "01:01:01"), (None, "00:00:00")],
)
def test_format_timestamp(seconds, expected) -> None:
    assert _format_timestamp(seconds) == expected


def test_whisper_format_result_with_chunks() -> None:
    result = {
        "text": "ignorado",
        "chunks": [
            {"timestamp": (51.0, 55.0), "text": " Combo e venda casada continuam valendo."},
            {"timestamp": (55.0, 58.0), "text": ""},
        ],
    }
    formatted = WhisperEngine._format_result(result)
    assert formatted == "[00:00:51] Combo e venda casada continuam valendo."


def test_whisper_format_result_without_chunks() -> None:
    assert WhisperEngine._format_result({"text": " olá mundo "}) == "olá mundo"
